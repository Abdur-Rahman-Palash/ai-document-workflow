from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel
from typing import Any, Dict, Optional
from io import BytesIO
import base64
import json

try:
    import fitz
except ImportError:
    fitz = None

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    Document = None

app = FastAPI(title="AI Document Engine")


class RecreateRequest(BaseModel):
    templateSchema: Dict[str, Any]
    structuredData: Dict[str, Any]
    outputFormat: str = "docx"


class UpdateDocumentRequest(BaseModel):
    contentBase64: str
    filename: str
    instructions: str
    outputFormat: str = "docx"
    updateSchema: Optional[Dict[str, Any]] = None


class GenerateDocumentRequest(BaseModel):
    templateSchema: Optional[Dict[str, Any]] = None
    structuredData: Dict[str, Any]
    outputFormat: str = "docx"


class ExportDocumentRequest(BaseModel):
    documentBase64: str
    outputFormat: str


@app.get("/health")
async def health() -> Dict[str, str]:
    return {"status": "healthy", "service": "ai-document-engine"}


@app.post("/analyze-document-style")
async def analyze_document_style(
    request: Request,
    file: UploadFile | None = File(None),
    filename: str | None = Form(None),
    contentBase64: str | None = Form(None),
) -> Dict[str, Any]:
    if file is not None:
        content = await file.read()
        filename = file.filename
    else:
        if request.headers.get("content-type", "").startswith("application/json"):
            body = await request.json()
            filename = body.get("filename")
            contentBase64 = body.get("contentBase64")
        if not filename or not contentBase64:
            raise HTTPException(status_code=400, detail="Missing filename or contentBase64")
        content = base64.b64decode(contentBase64)

    if not filename:
        raise HTTPException(status_code=400, detail="Missing filename")

    lower_name = filename.lower()
    if lower_name.endswith(".docx"):
        if Document is None:
            raise HTTPException(status_code=500, detail="python-docx not installed")
        doc = Document(BytesIO(content))
        return parse_docx_style(doc)
    if lower_name.endswith(".pdf"):
        if fitz is None:
            raise HTTPException(status_code=500, detail="PyMuPDF not installed")
        return parse_pdf_style(content)
    raise HTTPException(status_code=400, detail="Unsupported document type")


@app.post("/recreate-document")
async def recreate_document(request: RecreateRequest) -> Response:
    return build_docx_document(request.structuredData, request.outputFormat)


@app.post("/update-document-content")
async def update_document_content(request: UpdateDocumentRequest) -> Response:
    return build_docx_document({"instructions": request.instructions}, request.outputFormat)


@app.post("/generate-document-from-template")
async def generate_document_from_template(request: GenerateDocumentRequest) -> Response:
    return build_docx_document(request.structuredData, request.outputFormat)


@app.post("/export-document")
async def export_document(request: ExportDocumentRequest) -> Response:
    document_bytes = base64.b64decode(request.documentBase64)
    if request.outputFormat.lower() == "docx":
        return Response(content=document_bytes, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    if request.outputFormat.lower() == "pdf":
        if fitz is None:
            raise HTTPException(status_code=500, detail="PDF conversion requires PyMuPDF")
        return docx_to_pdf_response(document_bytes)
    raise HTTPException(status_code=400, detail="Unsupported export format")


def parse_docx_style(doc: "Document") -> Dict[str, Any]:
    fonts = {}
    sections = []
    for paragraph in doc.paragraphs:
        style = paragraph.style.name
        text = paragraph.text.strip()
        if text:
            fonts.setdefault(style, set()).add(paragraph.style.font.name or "Unknown")
            sections.append({"style": style, "text": text})
    return {
        "page": {"pageCount": len(doc.sections)},
        "fonts": {k: list(v) for k, v in fonts.items()},
        "sections": sections,
        "tables": [parse_table(table) for table in doc.tables],
        "spacing": {"default": "single"},
    }


def parse_table(table: Any) -> Dict[str, Any]:
    return {
        "rows": len(table.rows),
        "columns": len(table.columns),
        "headers": [cell.text for cell in table.rows[0].cells] if table.rows else [],
    }


def parse_pdf_style(content: bytes) -> Dict[str, Any]:
    doc = fitz.open(stream=content, filetype="pdf")
    metadata = {"pages": doc.page_count}
    fonts = set()
    for page in doc:
        for block in page.get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    fonts.add(span.get("font", "Unknown"))
    return {
        "page": {"pageCount": doc.page_count},
        "fonts": {"detected": list(fonts)},
        "sections": [{"page": idx + 1, "text": page.get_text("text").strip()} for idx, page in enumerate(doc)],
        "tables": [],
        "spacing": {"default": "unknown"},
    }


def build_docx_document(data: Dict[str, Any], output_format: str) -> Response:
    if Document is None:
        raise HTTPException(status_code=500, detail="python-docx not installed")
    doc = Document()
    doc.add_heading("Generated Document", level=1)
    doc.add_paragraph(json.dumps(data, indent=2))
    buffer = BytesIO()
    doc.save(buffer)
    payload = buffer.getvalue()
    if output_format.lower() == "docx":
        return Response(
            content=payload,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    if output_format.lower() == "pdf":
        return docx_to_pdf_response(payload)
    raise HTTPException(status_code=400, detail="Unsupported output format")


def docx_to_pdf_response(docx_bytes: bytes) -> Response:
    if fitz is None:
        raise HTTPException(status_code=500, detail="PDF conversion requires PyMuPDF")
    doc = fitz.open(stream=docx_bytes, filetype="docx")
    pdf_bytes = doc.convert_to_pdf()
    return Response(content=pdf_bytes, media_type="application/pdf")
